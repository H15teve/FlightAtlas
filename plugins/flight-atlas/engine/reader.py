"""Header-driven import. Ignore unrelated cells, formula execution and instructions."""
import csv,datetime as dt,json,re,zipfile,posixpath
from pathlib import Path
from xml.etree import ElementTree as ET
from models import normalize_model

HEADERS={
 'date':['日期','起飞日期','航班日期','出发日期','date','departure_date'],
 'flight':['航班号','航班','flight','flight_number'],
 'codeshare':['是否共享航班','共享航班','is_codeshare','codeshare'],
 'departure':['出发机场','出发城市','出发','起飞机场','出发地','departure','origin'],
 'arrival':['到达机场','到达城市','到达','降落机场','目的地','arrival','destination'],
 'departure_time':['实际起飞','实际起飞时间','起飞时间','出发时间','departure_time'],
 'arrival_time':['实际降落','实际到达时间','到达时间','降落时间','arrival_time'],
 'scheduled_departure':['表定出发','表定起飞','scheduled_departure'],
 'scheduled_arrival':['表定到达','scheduled_arrival'],
 'arrival_date':['到达日期','arrival_date'],
 'distance':['飞行里程','里程(公里)','里程','里程数','距离','飞行距离','distance','distance_km'],
 'registration':['飞机注册号','注册号','飞机编号','机号','registration','tail_number'],
 'model':['ICAO机型代码','ICAO机型','ICAO型号','机型ICAO','机型代码','机型','机型/子型号','型号','model','aircraft_type'],
 'carrier':['实际承运航司','实际承运航空','实际承运','执飞航司','实际执飞航司','实际运营航司','operating_airline','operator'],
 'marketing':['航空公司','航司','航空公司名称','营销航司','airline','carrier'],
 'delivery':['飞机交付日期','飞机交付时间','首次交付日期','首次交付时间','交付日期','交付时间','delivery_date'],
 'msn':['MSN','序列号/MSN','制造序列号','生产序列号','msn'],
 'status':['客票状态','航段状态','flight_status'],
 'minutes':['实际飞行时长','飞行时长(分钟)','飞行时长分钟','duration_minutes','minutes'],
 'scheduled_minutes':['表定飞行时长','scheduled_duration_minutes'],
}
NON_FLOWN={'未使用','未乘机','已退票','退票','已取消','取消','作废','void','cancelled','refunded','unused','open'}

def excluded(safe):
    if str(safe.get('flight','')).strip().upper().startswith('OPEN'):return 'OPEN excluded'
    if str(safe.get('status') or '').strip().lower() in NON_FLOWN:return 'explicit non-flown status excluded'
    return None
def label(x):return re.sub(r'[\s（）()_/-]','',str(x or '')).lower()
def date_text(x):
    if isinstance(x,(dt.datetime,dt.date)):return x.date().isoformat() if isinstance(x,dt.datetime) else x.isoformat()
    if isinstance(x,(float,int)):return (dt.datetime(1899,12,30)+dt.timedelta(days=x)).date().isoformat()
    text=str(x or '').strip();text=re.sub(r'[./年]', '-',text).replace('月','-').replace('日','')
    if 'T' in text:
        try:return dt.datetime.fromisoformat(text).date().isoformat()
        except ValueError:pass
    if re.fullmatch(r'\d{4}-\d{1,2}-\d{1,2}',text):return dt.date(*map(int,text.split('-'))).isoformat()
    return text
def registration(x):
    s=re.sub(r'\s','',str(x or '')).upper()
    if s in ['NONE','NAN','—','-']:return ''
    if re.fullmatch('B[A-Z0-9]{4}',s):s='B-'+s[1:]
    if s and not re.fullmatch('[A-Z0-9]+(?:-[A-Z0-9]+)?',s):raise ValueError('注册号格式无效')
    return s

def codeshare_flag(value):
    if value is None or str(value).strip() in ['', '-', '—', '--']:return None
    text=str(value).strip().lower()
    if text in ['是','共享','共享航班','yes','true','1','1.0']:return True
    if text in ['否','非共享','非共享航班','no','false','0','0.0']:return False
    raise ValueError('是否共享航班值无法识别；请用户核验，不默认当作非共享')

def _xlsx_cached_tables(path):
    """Read only declared worksheet cached values in malformed ZIP-based exports.

    Some Umetrip .xls files contain XLSX XML with invalid content-type metadata.
    No repair/write, macro execution, formula evaluation or external relationships.
    """
    from defusedxml import ElementTree
    from openpyxl.styles.numbers import BUILTIN_FORMATS,is_date_format
    from openpyxl.utils.datetime import from_excel,CALENDAR_MAC_1904,CALENDAR_WINDOWS_1900
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(path) as archive:
        if not {'xl/workbook.xml','xl/_rels/workbook.xml.rels'}<=set(archive.namelist()):raise ValueError('ZIP没有可识别的Excel工作簿结构')
        def xml(name):
            if archive.getinfo(name).file_size>100000000:raise ValueError('Excel XML过大，拒绝解析')
            return ElementTree.fromstring(archive.read(name))
        workbook=xml('xl/workbook.xml');properties=workbook.find('s:workbookPr',ns)
        epoch=CALENDAR_MAC_1904 if properties is not None and properties.get('date1904') in ['1','true'] else CALENDAR_WINDOWS_1900
        rels={e.get('Id'):e.get('Target') for e in xml('xl/_rels/workbook.xml.rels') if e.get('TargetMode')!='External'}
        strings=[]
        if 'xl/sharedStrings.xml' in archive.namelist():strings=[''.join(e.itertext()) for e in xml('xl/sharedStrings.xml')]
        date_styles=set()
        if 'xl/styles.xml' in archive.namelist():
            styles=xml('xl/styles.xml');formats={int(e.get('numFmtId')):e.get('formatCode','') for e in styles.findall('s:numFmts/s:numFmt',ns)}
            for i,e in enumerate(styles.findall('s:cellXfs/s:xf',ns)):
                number=int(e.get('numFmtId','0'))
                if is_date_format(formats.get(number,BUILTIN_FORMATS.get(number,''))):date_styles.add(i)
        sheets=sorted(workbook.findall('s:sheets/s:sheet',ns),key=lambda s:s.get('name')!='航段记录')
        for sheet in sheets:
            rid=sheet.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id');target=rels.get(rid,'')
            name=posixpath.normpath(target.lstrip('/') if target.startswith('/') else 'xl/'+target)
            if not name.startswith('xl/') or name not in archive.namelist():raise ValueError('工作表关系路径无效')
            table=[]
            for row in xml(name).findall('s:sheetData/s:row',ns):
                row_number=int(row.get('r',len(table)+1));table.extend([[] for _ in range(max(0,row_number-len(table)-1))]);values=[]
                for cell in row:
                    letters=re.sub('[^A-Z]','',cell.get('r','A'));index=0
                    for c in letters:index=index*26+ord(c)-64
                    if index<1 or index>16384:raise ValueError('工作表列号无效')
                    values.extend([None]*(index-len(values)))
                    value=cell.find('s:v',ns);text=value.text if value is not None else None;kind=cell.get('t')
                    if kind=='inlineStr':value=cell.find('s:is',ns);text=''.join(value.itertext()) if value is not None else ''
                    elif kind=='s' and text is not None:text=strings[int(text)]
                    elif text is not None and kind not in ['str','e','b']:
                        text=float(text)
                        if int(cell.get('s','0')) in date_styles:text=from_excel(text,epoch)
                    values[index-1]=text
                table.append(values)
            yield sheet.get('name','Sheet'),table
def _tables(path):
    ext=path.suffix.lower();head=path.open('rb').read(8)
    # BIFF/CFB files can contain an embedded ZIP theme stream. Header wins over
    # zipfile's end-of-file signature scan; otherwise valid Umetrip XLS is misread.
    if head!=b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' and zipfile.is_zipfile(path):
        import openpyxl
        fp=path.open('rb')
        try:wb=openpyxl.load_workbook(fp,read_only=True,data_only=True,keep_links=False)
        except (OSError,KeyError):
            fp.close();yield from _xlsx_cached_tables(path);return
        try:
            sheets=sorted(wb.worksheets,key=lambda s:s.title!='航段记录')
            for sheet in sheets:yield sheet.title,list(sheet.iter_rows(values_only=True))
        finally:wb.close();fp.close()
    elif head==b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
        import xlrd
        wb=xlrd.open_workbook(path,on_demand=True)
        try:
            for sheet in wb.sheets():
                values=[]
                for r in range(sheet.nrows):
                    row=[]
                    for c in range(sheet.ncols):
                        cell=sheet.cell(r,c);value=cell.value
                        if cell.ctype==xlrd.XL_CELL_DATE:value=xlrd.xldate_as_datetime(value,wb.datemode)
                        row.append(value)
                    values.append(row)
                yield sheet.name,values
        finally:wb.release_resources()
    elif ext in ['.csv','.tsv']:
        with path.open(encoding='utf-8-sig',newline='') as f:yield path.stem,list(csv.reader(f,delimiter='\t' if ext=='.tsv' else ','))
    elif b'<?xml' in head or head.startswith(b'\xef\xbb\xbf<'):
        from defusedxml import ElementTree
        tree=ElementTree.parse(path);ns={'s':'urn:schemas-microsoft-com:office:spreadsheet'}
        for sheet in tree.findall('.//s:Worksheet',ns):
            values=[]
            for r in sheet.findall('s:Table/s:Row',ns):
                row=[]
                for c in r:
                    idx=int(c.get('{'+ns['s']+'}Index',len(row)+1));row.extend(['']*(idx-len(row)-1));v=c.find('s:Data',ns);row.append(v.text if v is not None else '')
                values.append(row)
            yield sheet.get('{'+ns['s']+'}Name','Sheet'),values
    else:raise ValueError('仅支持真实xls、xlsx、Excel XML或UTF-8 CSV；不能按扩展名猜测内容。')

def read_flights(path,mapping=None):
    if Path(path).suffix.lower()=='.json':
        value=json.loads(Path(path).read_text(encoding='utf-8'))
        if not isinstance(value,dict) or value.get('kind')!='flightatlas-history' or value.get('version')!=1 or not isinstance(value.get('records'),list):
            raise ValueError('不是受支持的FlightAtlas历史文件')
        rows=[];warnings=[]
        for n,record in enumerate(value['records'],1):
            if not isinstance(record,dict):raise ValueError('历史航段需为对象')
            safe={k:v for k,v in record.items() if k in HEADERS or k in ['model_source','minutes_header','scheduled_minutes_header']}
            if any(isinstance(v,(dict,list)) for v in safe.values()):raise ValueError('历史航段字段需标量')
            reason=excluded(safe)
            if reason:warnings.append({'row':n,'reason':reason});continue
            if not all(safe.get(k) for k in ['flight','departure','arrival']):raise ValueError('历史航段缺少航班号或机场')
            safe['date']=date_text(safe.get('date'));safe['arrival_date']=date_text(safe.get('arrival_date'));safe['delivery']=date_text(safe.get('delivery'))
            safe['registration']=registration(safe.get('registration'));safe['model_source']=str(safe.get('model_source') or safe.get('model') or '')
            safe['model']=normalize_model(safe.get('model'));safe['codeshare']=codeshare_flag(safe.get('codeshare'))
            safe['carrier']=str(safe.get('carrier') or safe.get('marketing') or '未记录').strip();safe['row']=n
            rows.append(safe)
        fields={k:k for record in rows for k in record if k in HEADERS}
        return rows,{'schema':{'format':'flightatlas-history','version':1,'fields':fields},'warnings':warnings,'sources':value.get('sources',[]),'privacy':'Only whitelisted reporting fields retained'}
    mapping=mapping or {};rows=[];warnings=[];selected=None
    for sheet,table in _tables(Path(path)):
        if any(term in sheet.lower() for term in ['open票','无效票','退票','cancelled']):continue
        for h,raw in enumerate(table[:30]):
            labels=[label(x) for x in raw];found={}
            for key,aliases in HEADERS.items():
                names=[mapping[key]] if key in mapping else aliases
                indices=[labels.index(label(a)) for a in names if label(a) in labels]
                if indices:found[key]=indices[0]
            if not {'flight','departure','arrival'}<=found.keys():continue
            selected={'sheet':sheet,'header_row':h+1,'fields':{k:str(raw[v]) for k,v in found.items()}}
            for n,values in enumerate(table[h+1:],h+2):
                safe={k:values[i] if i<len(values) else None for k,i in found.items()}
                if not safe.get('flight') and not safe.get('departure') and not safe.get('arrival'):continue
                reason=excluded(safe)
                if reason:warnings.append({'row':n,'reason':reason});continue
                if not all(safe.get(k) for k in ['flight','departure','arrival']):raise ValueError(f'行{n}缺少航班号或机场；不能静默漏计。')
                safe['model_source']=str(safe.get('model') or '').strip()
                safe['codeshare']=codeshare_flag(safe.get('codeshare'))
                if 'minutes' in found:safe['minutes_header']=str(raw[found['minutes']])
                if 'scheduled_minutes' in found:safe['scheduled_minutes_header']=str(raw[found['scheduled_minutes']])
                safe.update({'row':n,'date':date_text(safe.get('date')),'arrival_date':date_text(safe.get('arrival_date')),'delivery':date_text(safe.get('delivery')),'registration':registration(safe.get('registration')),'model':normalize_model(safe.get('model'))})
                if 'PCF' in safe['model_source'].upper():warnings.append({'row':n,'reason':'原表机型含货运转换标记；仅归并型号，不作为历史执飞或退出客运证据'})
                if not safe.get('carrier'):warnings.append({'row':n,'reason':'承运人未提供，使用原表航司；未核验共享航班' if safe.get('marketing') else '承运航司字段未提供；保持未知，后续仅按显式配置识别'})
                safe['carrier']=str(safe.get('carrier') or safe.get('marketing') or '未记录').strip()
                rows.append(safe)
            return rows,{'schema':selected,'warnings':warnings,'model_normalization':[{'row':r['row'],'source':r['model_source'],'normalized':r['model']} for r in rows],'privacy':'Only whitelisted reporting fields retained; tickets, passenger data and workbook instructions discarded'}
    raise ValueError('无法识别航班表头。用配置column_map指定对应字段名称。')
